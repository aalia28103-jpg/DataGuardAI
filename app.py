import streamlit as st
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from openai import OpenAI

import io
import json
import re


# =========================================================
# PAGE CONFIG
# =========================================================

st.set_page_config(
    page_title="SlideGuard AI",
    page_icon="🛡️",
    layout="wide"
)


# =========================================================
# CUSTOM CSS
# =========================================================

st.markdown("""
<style>

.main-title {
    font-size: 42px;
    font-weight: 700;
    margin-bottom: 5px;
}

.subtitle {
    font-size: 18px;
    color: #666666;
    margin-bottom: 25px;
}

.section-title {
    font-size: 25px;
    font-weight: 600;
    margin-top: 25px;
}

</style>
""", unsafe_allow_html=True)


# =========================================================
# THEME COLOURS
# =========================================================

THEME_COLORS = {

    "🔵 Blue": {
        "main": RGBColor(31, 78, 121),
        "dark": RGBColor(20, 50, 80)
    },

    "🟢 Green": {
        "main": RGBColor(46, 125, 50),
        "dark": RGBColor(30, 85, 35)
    },

    "🟣 Purple": {
        "main": RGBColor(106, 76, 147),
        "dark": RGBColor(70, 45, 100)
    },

    "🟠 Orange": {
        "main": RGBColor(230, 126, 34),
        "dark": RGBColor(150, 75, 15)
    },

    "🔴 Red": {
        "main": RGBColor(192, 57, 43),
        "dark": RGBColor(125, 35, 25)
    },

    "⚫ Dark": {
        "main": RGBColor(45, 45, 45),
        "dark": RGBColor(20, 20, 20)
    },

    "🔷 Teal": {
        "main": RGBColor(0, 121, 140),
        "dark": RGBColor(0, 75, 90)
    }
}


# =========================================================
# BACKGROUND COLOURS
# =========================================================

BACKGROUND_COLORS = {

    "⚪ White": RGBColor(255, 255, 255),

    "🔵 Very Light Blue": RGBColor(235, 245, 255),

    "🟢 Very Light Green": RGBColor(237, 248, 237),

    "🟣 Very Light Purple": RGBColor(245, 239, 250),

    "🟠 Very Light Orange": RGBColor(255, 245, 232),

    "🔴 Very Light Red": RGBColor(255, 240, 238),

    "🔷 Very Light Teal": RGBColor(235, 249, 250),

    "⚫ Light Gray": RGBColor(242, 242, 242)
}


# =========================================================
# FONT OPTIONS
# =========================================================

FONT_OPTIONS = [
    "Aptos",
    "Arial",
    "Calibri",
    "Times New Roman",
    "Georgia",
    "Verdana",
    "Tahoma",
    "Trebuchet MS",
    "Courier New"
]


# =========================================================
# SESSION STATE
# =========================================================

if "slides" not in st.session_state:
    st.session_state.slides = []

if "generated_ppt" not in st.session_state:
    st.session_state.generated_ppt = None


# =========================================================
# CLEAN TEXT
# =========================================================

def clean_text(text):

    if not text:
        return ""

    text = text.replace("\r", "\n")

    return text.strip()


# =========================================================
# OPENAI CLIENT
# =========================================================

def get_openai_client():

    try:

        api_key = st.secrets["OPENAI_API_KEY"]

        if not api_key:
            return None

        return OpenAI(api_key=api_key)

    except Exception:

        return None


# =========================================================
# AI TEXT ENHANCEMENT
# =========================================================

def enhance_slide_with_ai(title, content):

    client = get_openai_client()

    if client is None:
        return title, content

    prompt = f"""
You are a professional presentation creator.

Convert the following slide content into concise PowerPoint-friendly
bullet points.

Slide title:
{title}

Slide content:
{content}

Rules:
- Keep the original meaning.
- Do not invent facts.
- Create 3 to 6 concise bullet points.
- Each bullet should be easy to read on a presentation slide.
- Do not use markdown.
- Return ONLY valid JSON.

Format:
{{
    "title": "short slide title",
    "bullets": [
        "bullet 1",
        "bullet 2",
        "bullet 3"
    ]
}}
"""

    try:

        response = client.responses.create(
            model="gpt-6-luna",
            input=prompt
        )

        result = response.output_text.strip()

        result = re.sub(
            r"```json|```",
            "",
            result
        ).strip()

        data = json.loads(result)

        new_title = data.get("title", title)

        bullets = data.get("bullets", [])

        if isinstance(bullets, list):

            new_content = "\n".join(
                f"• {str(b).strip()}"
                for b in bullets
                if str(b).strip()
            )

            return new_title, new_content

    except Exception:

        st.warning(
            "AI enhancement could not be applied. "
            "Original content will be used."
        )

    return title, content


# =========================================================
# ADD TEXT BOX
# =========================================================

def add_text_box(
    slide,
    text,
    left,
    top,
    width,
    height,
    font_name="Aptos",
    font_size=20,
    color=RGBColor(40, 40, 40),
    bold=False,
    italic=False,
    alignment=PP_ALIGN.LEFT
):

    box = slide.shapes.add_textbox(
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height)
    )

    tf = box.text_frame

    tf.clear()

    tf.word_wrap = True

    tf.vertical_anchor = MSO_ANCHOR.TOP

    p = tf.paragraphs[0]

    p.text = text

    p.alignment = alignment

    p.font.name = font_name

    p.font.size = Pt(font_size)

    p.font.bold = bold

    p.font.italic = italic

    p.font.color.rgb = color

    return box


# =========================================================
# ADD FOOTER
# =========================================================

def add_footer(
    slide,
    theme,
    font_name
):

    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0),
        Inches(7.25),
        Inches(13.333),
        Inches(0.08)
    )

    line.fill.solid()

    line.fill.fore_color.rgb = theme["main"]

    line.line.fill.background()

    add_text_box(
        slide,
        "SlideGuard AI | Developed by Aliya Banu A",
        0.45,
        7.32,
        12.4,
        0.25,
        font_name=font_name,
        font_size=9,
        color=RGBColor(120, 120, 120),
        alignment=PP_ALIGN.RIGHT
    )


# =========================================================
# TITLE SLIDE
# =========================================================

def create_title_slide(
    prs,
    title,
    presented_by,
    theme,
    background_color,
    font_name,
    font_size,
    text_color,
    bold,
    italic,
    alignment
):

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    # Background
    background = slide.background

    fill = background.fill

    fill.solid()

    fill.fore_color.rgb = background_color


    # Top colour block
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0),
        Inches(0),
        Inches(13.333),
        Inches(1.0)
    )

    shape.fill.solid()

    shape.fill.fore_color.rgb = theme["main"]

    shape.line.fill.background()


    # Presentation title
    add_text_box(
        slide,
        title,
        0.8,
        2.2,
        11.7,
        1.4,
        font_name=font_name,
        font_size=font_size + 10,
        color=theme["dark"],
        bold=True,
        italic=italic,
        alignment=PP_ALIGN.CENTER
    )


    # Presented By
    if presented_by:

        add_text_box(
            slide,
            f"Presented By: {presented_by}",
            1.2,
            3.7,
            10.9,
            0.8,
            font_name=font_name,
            font_size=font_size,
            color=text_color,
            bold=bold,
            italic=italic,
            alignment=PP_ALIGN.CENTER
        )


    # Accent
    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(4.5),
        Inches(5.2),
        Inches(4.3),
        Inches(0.12)
    )

    accent.fill.solid()

    accent.fill.fore_color.rgb = theme["main"]

    accent.line.fill.background()


    add_footer(
        slide,
        theme,
        font_name
    )


# =========================================================
# CONTENT SLIDE
# =========================================================

def create_content_slide(
    prs,
    title,
    content,
    image_bytes,
    theme,
    background_color,
    font_name,
    font_size,
    text_color,
    bold,
    italic,
    alignment
):

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )


    # Background
    background = slide.background

    fill = background.fill

    fill.solid()

    fill.fore_color.rgb = background_color


    # Header
    header = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0),
        Inches(0),
        Inches(13.333),
        Inches(0.75)
    )

    header.fill.solid()

    header.fill.fore_color.rgb = theme["main"]

    header.line.fill.background()


    # Slide title
    add_text_box(
        slide,
        title,
        0.5,
        0.13,
        12.2,
        0.45,
        font_name=font_name,
        font_size=25,
        color=RGBColor(255, 255, 255),
        bold=True
    )


    # =====================================================
    # CONTENT WITH IMAGE
    # =====================================================

    if image_bytes:

        content_box = slide.shapes.add_textbox(
            Inches(0.65),
            Inches(1.15),
            Inches(6.8),
            Inches(5.65)
        )

        tf = content_box.text_frame

        tf.clear()

        tf.word_wrap = True

        lines = content.split("\n")

        first = True

        for line in lines:

            line = line.strip()

            if not line:
                continue

            if first:

                p = tf.paragraphs[0]

                first = False

            else:

                p = tf.add_paragraph()

            p.text = line

            p.font.name = font_name

            p.font.size = Pt(font_size)

            p.font.color.rgb = text_color

            p.font.bold = bold

            p.font.italic = italic

            p.alignment = alignment

            p.space_after = Pt(10)


        # Image
        try:

            image_stream = io.BytesIO(
                image_bytes
            )

            slide.shapes.add_picture(
                image_stream,
                Inches(7.85),
                Inches(1.35),
                width=Inches(4.85),
                height=Inches(4.85)
            )

        except Exception:

            pass


    # =====================================================
    # CONTENT WITHOUT IMAGE
    # =====================================================

    else:

        content_box = slide.shapes.add_textbox(
            Inches(0.8),
            Inches(1.25),
            Inches(11.8),
            Inches(5.5)
        )

        tf = content_box.text_frame

        tf.clear()

        tf.word_wrap = True

        lines = content.split("\n")

        first = True

        for line in lines:

            line = line.strip()

            if not line:
                continue

            if first:

                p = tf.paragraphs[0]

                first = False

            else:

                p = tf.add_paragraph()

            p.text = line

            p.font.name = font_name

            p.font.size = Pt(font_size)

            p.font.color.rgb = text_color

            p.font.bold = bold

            p.font.italic = italic

            p.alignment = alignment

            p.space_after = Pt(12)


    add_footer(
        slide,
        theme,
        font_name
    )


# =========================================================
# CREATE POWERPOINT
# =========================================================

def create_presentation(
    presentation_title,
    presented_by,
    slides_data,
    theme,
    background_color,
    font_name,
    font_size,
    text_color,
    bold,
    italic,
    alignment
):

    prs = Presentation()

    # 16:9 widescreen
    prs.slide_width = Inches(13.333)

    prs.slide_height = Inches(7.5)


    # Title slide
    create_title_slide(
        prs,
        presentation_title,
        presented_by,
        theme,
        background_color,
        font_name,
        font_size,
        text_color,
        bold,
        italic,
        alignment
    )


    # Content slides
    for slide_data in slides_data:

        create_content_slide(
            prs,
            slide_data["title"],
            slide_data["content"],
            slide_data["image"],
            theme,
            background_color,
            font_name,
            font_size,
            text_color,
            bold,
            italic,
            alignment
        )


    output = io.BytesIO()

    prs.save(output)

    output.seek(0)

    return output


# =========================================================
# HEADER
# =========================================================

st.markdown(
    '<div class="main-title">🛡️ SlideGuard AI</div>',
    unsafe_allow_html=True
)

st.markdown(
    '<div class="subtitle">'
    'AI-Powered Presentation Generator'
    '</div>',
    unsafe_allow_html=True
)


# =========================================================
# SIDEBAR
# =========================================================

st.sidebar.title(
    "🎨 Presentation Settings"
)


# Theme
theme_name = st.sidebar.selectbox(
    "🎨 Theme Colour",
    list(THEME_COLORS.keys())
)

theme = THEME_COLORS[theme_name]


# Background
background_name = st.sidebar.selectbox(
    "🖼️ Background Colour",
    list(BACKGROUND_COLORS.keys())
)

background_color = BACKGROUND_COLORS[
    background_name
]


st.sidebar.markdown("---")


# =========================================================
# TEXT FORMATTING
# =========================================================

st.sidebar.subheader(
    "✍️ Text Formatting"
)


# Font
font_name = st.sidebar.selectbox(
    "Font",
    FONT_OPTIONS
)


# Font size
font_size = st.sidebar.slider(
    "Font Size",
    min_value=12,
    max_value=32,
    value=20,
    step=1
)


# Font style
font_style = st.sidebar.selectbox(
    "Font Style",
    [
        "Normal",
        "Bold",
        "Italic",
        "Bold Italic"
    ]
)


bold = font_style in [
    "Bold",
    "Bold Italic"
]

italic = font_style in [
    "Italic",
    "Bold Italic"
]


# Text colour
text_color_name = st.sidebar.selectbox(
    "Text Colour",
    [
        "⚫ Black",
        "🔵 Dark Blue",
        "🟢 Dark Green",
        "🟣 Dark Purple",
        "🔴 Dark Red",
        "⚪ White"
    ]
)


TEXT_COLORS = {

    "⚫ Black": RGBColor(40, 40, 40),

    "🔵 Dark Blue": RGBColor(20, 50, 80),

    "🟢 Dark Green": RGBColor(30, 85, 35),

    "🟣 Dark Purple": RGBColor(70, 45, 100),

    "🔴 Dark Red": RGBColor(125, 35, 25),

    "⚪ White": RGBColor(255, 255, 255)
}


text_color = TEXT_COLORS[
    text_color_name
]


# Alignment
alignment_name = st.sidebar.selectbox(
    "📐 Text Alignment",
    [
        "⬅️ Left",
        "↔️ Center",
        "➡️ Right",
        "📏 Justify"
    ]
)


if alignment_name == "⬅️ Left":

    alignment = PP_ALIGN.LEFT

elif alignment_name == "↔️ Center":

    alignment = PP_ALIGN.CENTER

elif alignment_name == "➡️ Right":

    alignment = PP_ALIGN.RIGHT

else:

    alignment = PP_ALIGN.JUSTIFY

st.sidebar.markdown("---")


st.sidebar.info(
    "Choose the theme, background colour and "
    "text formatting before generating your presentation."
)


# =========================================================
# PRESENTATION DETAILS
# =========================================================

st.markdown(
    '<div class="section-title">'
    '📋 Presentation Details'
    '</div>',
    unsafe_allow_html=True
)


presentation_title = st.text_input(
    "Presentation Title",
    placeholder=(
        "Example: Artificial Intelligence in Business"
    )
)


presented_by = st.text_input(
    "Presented By",
    placeholder="Example: Aliya Banu A"
)


number_of_slides = st.number_input(
    "Number of Content Slides",
    min_value=1,
    max_value=30,
    value=3,
    step=1
)


# =========================================================
# GENERATE SLIDE INPUTS
# =========================================================

if st.button(
    "📝 Generate Slide Inputs",
    use_container_width=True
):

    st.session_state.slides = []

    st.session_state.generated_ppt = None

    for i in range(
        int(number_of_slides)
    ):

        st.session_state.slides.append({

            "title": "",

            "content": "",

            "image": None
        })

    st.rerun()


# =========================================================
# SLIDE INPUTS
# =========================================================

if st.session_state.slides:

    st.markdown(
        '<div class="section-title">'
        '📝 Slide Content'
        '</div>',
        unsafe_allow_html=True
    )


    for i in range(
        len(st.session_state.slides)
    ):

        st.markdown(
            f"### 📌 Slide {i + 1}"
        )


        # Slide title
        slide_title = st.text_input(
            f"Slide {i + 1} Title",
            value=st.session_state.slides[i]["title"],
            key=f"title_{i}",
            placeholder="Enter slide title"
        )


        # Slide content
        slide_content = st.text_area(
            f"Slide {i + 1} Content",
            value=st.session_state.slides[i]["content"],
            key=f"content_{i}",
            height=150,
            placeholder=(
                "Paste your slide content here..."
            )
        )


        # =================================================
        # PICTURE UPLOAD
        # =================================================

        st.markdown(
            "🖼️ **Picture for this slide**"
        )

        image_file = st.file_uploader(
            f"Upload Picture for Slide {i + 1} "
            "(Optional)",
            type=[
                "png",
                "jpg",
                "jpeg"
            ],
            key=f"image_{i}"
        )


        if image_file:

            st.image(
                image_file,
                caption=f"Slide {i + 1} Picture",
                width=350
            )


        # Save slide data
        st.session_state.slides[i][
            "title"
        ] = slide_title

        st.session_state.slides[i][
            "content"
        ] = slide_content


        if image_file:

            st.session_state.slides[i][
                "image"
            ] = image_file.getvalue()


        st.divider()


# =========================================================
# AI TEXT ENHANCEMENT
# =========================================================

if st.session_state.slides:

    st.markdown(
        '<div class="section-title">'
        '✨ AI Enhancement'
        '</div>',
        unsafe_allow_html=True
    )


    ai_enabled = st.checkbox(
        "Use AI to improve slide content",
        value=False
    )


    st.caption(
        "AI converts long paragraphs into concise "
        "presentation-friendly bullet points."
    )


# =========================================================
# GENERATE POWERPOINT
# =========================================================

if st.session_state.slides:

    if st.button(
        "🚀 Generate PowerPoint",
        type="primary",
        use_container_width=True
    ):

        if not presentation_title.strip():

            st.error(
                "Please enter a presentation title."
            )

        else:

            valid = True


            for i, slide in enumerate(
                st.session_state.slides
            ):

                if not slide["title"].strip():

                    st.error(
                        f"Please enter a title "
                        f"for Slide {i + 1}."
                    )

                    valid = False


                if not slide["content"].strip():

                    st.error(
                        f"Please enter content "
                        f"for Slide {i + 1}."
                    )

                    valid = False


            if valid:

                processed_slides = []


                progress = st.progress(0)


                for i, slide in enumerate(
                    st.session_state.slides
                ):

                    title = slide["title"]

                    content = slide["content"]


                    # AI text enhancement
                    if ai_enabled:

                        title, content = (
                            enhance_slide_with_ai(
                                title,
                                content
                            )
                        )


                    processed_slides.append({

                        "title": title,

                        "content": content,

                        "image": slide["image"]
                    })


                    progress.progress(
                        (i + 1) /
                        len(
                            st.session_state.slides
                        )
                    )


                # Create PPT
                ppt = create_presentation(

                    presentation_title,

                    presented_by,

                    processed_slides,

                    theme,

                    background_color,

                    font_name,

                    font_size,

                    text_color,

                    bold,

                    italic,

                    alignment
                )


                st.session_state.generated_ppt = (
                    ppt.getvalue()
                )


                st.success(
                    "🎉 PowerPoint generated successfully!"
                )


# =========================================================
# DOWNLOAD
# =========================================================

if st.session_state.generated_ppt:

    st.markdown("---")

    st.subheader(
        "📥 Download Presentation"
    )


    st.download_button(

        label="📥 Download PowerPoint",

        data=st.session_state.generated_ppt,

        file_name=(
            "SlideGuard_AI_Presentation.pptx"
        ),

        mime=(
            "application/vnd.openxmlformats-officedocument."
            "presentationml.presentation"
        ),

        use_container_width=True
    )


# =========================================================
# FOOTER
# =========================================================

st.markdown("---")

st.caption(
    "🛡️ SlideGuard AI | Developed by Aliya Banu A"
)